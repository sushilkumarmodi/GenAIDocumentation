from docx import Document

def create_docx(file_name,content):
    doc = Document()
    doc.add_heading('Document Title', level=1)
    doc.add_paragraph(content)
    
    # Save the document
    doc.save(file_name)